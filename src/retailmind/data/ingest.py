"""Read every workbook sheet, preserve row lineage, and classify transactions."""

from __future__ import annotations

import hashlib
import json
import os
from datetime import date, datetime
from importlib.metadata import version
from pathlib import Path
from typing import Any

import duckdb
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from retailmind.config import ProjectConfig


SOURCE_URL = "https://archive.ics.uci.edu/dataset/502/online+retail+ii"
SOURCE_LICENSE = "CC BY 4.0"
COLUMNS = (
    "invoice",
    "stock_code",
    "description",
    "quantity",
    "invoice_date",
    "price",
    "customer_id",
    "country",
)
HEADER_ALIASES = {
    "invoice": "invoice",
    "invoiceno": "invoice",
    "stockcode": "stock_code",
    "description": "description",
    "quantity": "quantity",
    "invoicedate": "invoice_date",
    "price": "price",
    "unitprice": "price",
    "customerid": "customer_id",
    "country": "country",
}
RAW_SCHEMA = pa.schema(
    [
        ("raw_row_id", pa.string()),
        ("source_hash", pa.string()),
        ("source_sheet", pa.string()),
        ("source_row", pa.int32()),
        ("invoice_raw", pa.string()),
        ("stock_code_raw", pa.string()),
        ("description_raw", pa.string()),
        ("quantity_raw", pa.string()),
        ("invoice_date_raw", pa.string()),
        ("price_raw", pa.string()),
        ("customer_id_raw", pa.string()),
        ("country_raw", pa.string()),
    ]
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value)


def _canonical_header(value: Any) -> str:
    if value is None:
        return ""
    return "".join(char for char in str(value).lower() if char.isalnum())


def _column_positions(header: tuple[Any, ...]) -> dict[str, int]:
    positions: dict[str, int] = {}
    for index, value in enumerate(header):
        mapped = HEADER_ALIASES.get(_canonical_header(value))
        if mapped:
            if mapped in positions:
                raise ValueError(f"Duplicate mapped workbook header: {mapped}")
            positions[mapped] = index
    missing = set(COLUMNS) - positions.keys()
    if missing:
        raise ValueError(f"Missing workbook columns: {sorted(missing)}")
    return positions


def inspect_workbook(config: ProjectConfig) -> dict[str, Any]:
    """Inspect workbook metadata and verify both sheets' mapped schemas."""
    if not config.workbook.is_file():
        raise FileNotFoundError(f"Workbook not found: {config.workbook}")
    source_hash = sha256_file(config.workbook)
    workbook = load_workbook(config.workbook, read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            header = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
            positions = _column_positions(header)
            sheets.append(
                {
                    "name": sheet.title,
                    "declared_data_rows": sheet.max_row - 1,
                    "headers": [_text(value) for value in header],
                    "column_positions": positions,
                }
            )
    finally:
        workbook.close()
    result = {
        "source": SOURCE_URL,
        "license": SOURCE_LICENSE,
        "source_filename": config.workbook.name,
        "workbook_sha256": source_hash,
        "workbook_bytes": config.workbook.stat().st_size,
        "sheets": sheets,
        "declared_total_rows": sum(sheet["declared_data_rows"] for sheet in sheets),
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return result


def _sql_path(path: Path) -> str:
    return str(path.resolve()).replace("'", "''").replace("\\", "/")


def _write_raw(config: ProjectConfig, source_hash: str) -> list[dict[str, Any]]:
    config.processed_dir.mkdir(parents=True, exist_ok=True)
    destination = config.processed_dir / "raw_transactions.parquet"
    temporary = destination.with_suffix(".parquet.tmp")
    writer = pq.ParquetWriter(temporary, RAW_SCHEMA, compression="zstd")
    workbook = load_workbook(config.workbook, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    batch: list[dict[str, Any]] = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            iterator = sheet.iter_rows(values_only=True)
            header = tuple(next(iterator))
            positions = _column_positions(header)
            count = 0
            for row_number, row in enumerate(iterator, start=2):
                fields = {
                    column: _text(row[positions[column]] if positions[column] < len(row) else None)
                    for column in COLUMNS
                }
                batch.append(
                    {
                        "raw_row_id": f"{source_hash[:16]}:{sheet_index}:{row_number}",
                        "source_hash": source_hash,
                        "source_sheet": sheet.title,
                        "source_row": row_number,
                        "invoice_raw": fields["invoice"],
                        "stock_code_raw": fields["stock_code"],
                        "description_raw": fields["description"],
                        "quantity_raw": fields["quantity"],
                        "invoice_date_raw": fields["invoice_date"],
                        "price_raw": fields["price"],
                        "customer_id_raw": fields["customer_id"],
                        "country_raw": fields["country"],
                    }
                )
                count += 1
                if len(batch) >= 25_000:
                    writer.write_table(pa.Table.from_pylist(batch, schema=RAW_SCHEMA))
                    batch.clear()
            sheets.append({"name": sheet.title, "rows": count, "headers": [_text(v) for v in header]})
        if batch:
            writer.write_table(pa.Table.from_pylist(batch, schema=RAW_SCHEMA))
            batch.clear()
    finally:
        writer.close()
        workbook.close()
    os.replace(temporary, destination)
    return sheets


def _classify(config: ProjectConfig) -> dict[str, Any]:
    raw_path = _sql_path(config.processed_dir / "raw_transactions.parquet")
    classified_path = config.processed_dir / "transactions.parquet"
    temp_path = classified_path.with_suffix(".parquet.tmp")
    excluded = config.excluded_stock_codes
    exclusion_sql = (
        "stock_code IN (" + ", ".join("'" + code.replace("'", "''") + "'" for code in sorted(excluded)) + ")"
        if excluded
        else "FALSE"
    )
    connection = duckdb.connect()
    try:
        connection.execute(f"CREATE VIEW raw AS SELECT * FROM read_parquet('{raw_path}')")
        connection.execute(
            """
            CREATE VIEW parsed AS
            SELECT *,
                NULLIF(TRIM(invoice_raw), '') AS invoice_no,
                NULLIF(TRIM(stock_code_raw), '') AS stock_code,
                NULLIF(TRIM(description_raw), '') AS description,
                TRY_CAST(quantity_raw AS BIGINT) AS quantity,
                TRY_CAST(invoice_date_raw AS TIMESTAMP) AS invoice_date,
                TRY_CAST(price_raw AS DECIMAL(18, 4)) AS unit_price,
                NULLIF(REGEXP_REPLACE(TRIM(customer_id_raw), '\\.0$', ''), '') AS customer_id,
                NULLIF(TRIM(country_raw), '') AS country
            FROM raw
            """
        )
        connection.execute(
            f"""
            CREATE VIEW classified AS
            SELECT *,
                CASE
                    WHEN invoice_date IS NULL THEN 'invalid_date'
                    WHEN invoice_no IS NULL THEN 'missing_invoice'
                    WHEN stock_code IS NULL THEN 'missing_stock_code'
                    WHEN quantity IS NULL THEN 'invalid_quantity'
                    WHEN unit_price IS NULL THEN 'invalid_price'
                    WHEN LOWER(LEFT(invoice_no, 1)) = 'c' THEN 'cancellation'
                    WHEN quantity <= 0 THEN 'nonpositive_quantity'
                    WHEN unit_price <= 0 THEN 'nonpositive_price'
                    WHEN {exclusion_sql} THEN 'non_merchandise'
                    WHEN customer_id IS NULL THEN 'sale_anonymous'
                    ELSE 'sale_identified'
                END AS reason_code,
                quantity * unit_price AS line_value
            FROM parsed
            """
        )
        destination_sql = _sql_path(temp_path)
        connection.execute(
            f"COPY (SELECT * FROM classified) TO '{destination_sql}' "
            "(FORMAT PARQUET, COMPRESSION ZSTD)"
        )
        counts = dict(
            connection.execute(
                "SELECT reason_code, COUNT(*) FROM classified GROUP BY reason_code ORDER BY reason_code"
            ).fetchall()
        )
        total, unique_rows, min_date, max_date = connection.execute(
            "SELECT COUNT(*), COUNT(DISTINCT raw_row_id), MIN(invoice_date), MAX(invoice_date) "
            "FROM classified"
        ).fetchone()
        duplicate_extras = connection.execute(
            """
            SELECT COALESCE(SUM(n - 1), 0)
            FROM (
                SELECT COUNT(*) AS n FROM classified
                GROUP BY invoice_no, stock_code, description, quantity,
                         invoice_date, unit_price, customer_id, country
                HAVING COUNT(*) > 1
            )
            """
        ).fetchone()[0]
        raw_fields = (
            "invoice_raw", "stock_code_raw", "description_raw", "quantity_raw",
            "invoice_date_raw", "price_raw", "customer_id_raw", "country_raw",
        )
        missing_sql = ", ".join(
            f"COUNT(*) FILTER (WHERE NULLIF(TRIM({field}), '') IS NULL) AS {field}"
            for field in raw_fields
        )
        missing_values = dict(zip(
            raw_fields,
            connection.execute(f"SELECT {missing_sql} FROM raw").fetchone(),
            strict=True,
        ))
        sample_columns = (
            "reason_code", "raw_row_id", "source_sheet", "source_row",
            "invoice_raw", "stock_code_raw", "customer_id_raw",
        )
        sample_rows = connection.execute(
            """
            SELECT reason_code, raw_row_id, source_sheet, source_row,
                   invoice_raw, stock_code_raw, customer_id_raw
            FROM (
                SELECT *, ROW_NUMBER() OVER (
                    PARTITION BY reason_code ORDER BY raw_row_id
                ) AS sample_rank
                FROM classified
            )
            WHERE sample_rank <= 3
            ORDER BY reason_code, raw_row_id
            """
        ).fetchall()
        reason_samples: dict[str, list[dict[str, Any]]] = {}
        for row in sample_rows:
            example = dict(zip(sample_columns, row, strict=True))
            reason_samples.setdefault(example["reason_code"], []).append(example)
    finally:
        connection.close()
    os.replace(temp_path, classified_path)
    return {
        "total_rows": total,
        "unique_raw_row_ids": unique_rows,
        "reason_counts": counts,
        "missing_raw_values": missing_values,
        "reason_samples": reason_samples,
        "suspicious_duplicate_extra_rows": duplicate_extras,
        "min_valid_invoice_date": min_date.isoformat(sep=" ") if min_date else None,
        "max_valid_invoice_date": max_date.isoformat(sep=" ") if max_date else None,
        "reconciled": sum(counts.values()) == total == unique_rows,
    }


def prepare_workbook(config: ProjectConfig) -> dict[str, Any]:
    """Build deterministic Parquet outputs and reconciled public reports."""
    inspection = inspect_workbook(config)
    sheets = _write_raw(config, inspection["workbook_sha256"])
    quality = _classify(config)
    observed_rows = sum(sheet["rows"] for sheet in sheets)
    if observed_rows != quality["total_rows"] or not quality["reconciled"]:
        raise ValueError("Raw and classified row counts do not reconcile")
    config.reports_dir.mkdir(parents=True, exist_ok=True)
    manifest = {
        "source": SOURCE_URL,
        "license": SOURCE_LICENSE,
        "source_filename": config.workbook.name,
        "classification_version": "v1-exact-stock-code-exclusions",
        "excluded_stock_codes": sorted(config.excluded_stock_codes),
        "workbook_sha256": inspection["workbook_sha256"],
        "workbook_bytes": inspection["workbook_bytes"],
        "sheets": sheets,
        "total_rows": observed_rows,
        "raw_row_id_rule": "first 16 SHA-256 hex characters:zero-based sheet index:one-based Excel row",
        "outputs": ["data/processed/raw_transactions.parquet", "data/processed/transactions.parquet"],
        "package_versions": {
            package: version(package) for package in ("duckdb", "openpyxl", "pyarrow", "pyyaml")
        },
    }
    (config.reports_dir / "data_manifest.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    (config.reports_dir / "data_quality.json").write_text(
        json.dumps(quality, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(json.dumps({"manifest": manifest, "quality": quality}, indent=2))
    return {"manifest": manifest, "quality": quality}
